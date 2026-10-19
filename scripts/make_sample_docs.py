"""Generate synthetic Word documents with planted PII, plus gold sidecars.

    uv run python scripts/make_sample_docs.py                 # writes samples/
    uv run python scripts/make_sample_docs.py --out DIR --seed 7

Each ``<name>.docx`` gets a ``<name>.gold.json``: a list of
``{"value", "label", "location"}`` records. Locations are the engine's segment
locations (header, body, table, textbox, footer, comment) plus ``deleted``
(only inside a tracked deletion), ``comment_author`` and ``metadata`` (core
properties and revision authors). Values are written exactly as they appear.
"""

from __future__ import annotations

import argparse
import json
import random
from collections.abc import Callable
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape, quoteattr

import docx
from docx.document import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.text.paragraph import Paragraph
from faker import Faker

DEFAULT_OUT = Path(__file__).resolve().parents[1] / "samples"
VML_NS = "urn:schemas-microsoft-com:vml"
REVISION_DATE = "2025-10-14T09:30:00Z"
TFN_WEIGHTS = (1, 4, 3, 7, 5, 8, 6, 9, 10)
MEDICARE_WEIGHTS = (1, 3, 7, 9, 1, 3, 7, 9)
# Person names that would collide with hard negatives or ordinary words.
BLOCKED_NAMES = {
    "may", "june", "april", "august", "sydney", "opera", "house", "will", "mark",
    "grace", "hope", "rose", "dear", "victoria", "jordan", "street", "invoice",
}

Gold = list[dict[str, str]]


# --------------------------------------------------------------------------- #
# Checksummed identifiers
# --------------------------------------------------------------------------- #
def tfn_is_valid(number: str) -> bool:
    """Australian Tax File Number: weighted digit sum divisible by 11."""
    digits = [int(c) for c in number if c.isdigit()]
    return len(digits) == 9 and sum(d * w for d, w in zip(digits, TFN_WEIGHTS)) % 11 == 0


def medicare_is_valid(number: str) -> bool:
    """Medicare card number: first digit 2-6, 9th digit is the weighted check digit."""
    digits = [int(c) for c in number if c.isdigit()]
    return (
        len(digits) in (10, 11)
        and 2 <= digits[0] <= 6
        and sum(d * w for d, w in zip(digits, MEDICARE_WEIGHTS)) % 10 == digits[8]
    )


def make_tfn(rng: random.Random) -> str:
    while True:
        digits = [rng.randint(0, 9) for _ in range(8)]
        # (partial + 10 * d) % 11 == 0  <=>  d == partial % 11, since 10 == -1 (mod 11)
        check = sum(d * w for d, w in zip(digits, TFN_WEIGHTS)) % 11
        if check <= 9:
            s = "".join(map(str, [*digits, check]))
            return f"{s[:3]} {s[3:6]} {s[6:]}"


def make_medicare(rng: random.Random) -> str:
    digits = [rng.randint(2, 6)] + [rng.randint(0, 9) for _ in range(7)]
    digits.append(sum(d * w for d, w in zip(digits, MEDICARE_WEIGHTS)) % 10)
    digits.append(rng.randint(1, 9))  # card issue number
    s = "".join(map(str, digits))
    return f"{s[:4]} {s[4:9]} {s[9]}"


def make_mobile(rng: random.Random) -> str:
    return f"04{rng.randint(10, 99)} {rng.randint(100, 999)} {rng.randint(100, 999)}"


def make_landline(rng: random.Random) -> str:
    return f"(02) 9{rng.randint(100, 999)} {rng.randint(1000, 9999)}"


def make_bsb(rng: random.Random) -> str:
    return f"{rng.choice(['062', '033', '082', '012'])}-{rng.randint(0, 999):03d}"


# --------------------------------------------------------------------------- #
# People and gold records
# --------------------------------------------------------------------------- #
class People:
    """Draws distinct person names that cannot collide with the boilerplate."""

    def __init__(self, fake: Faker) -> None:
        self.fake = fake
        self.used: set[str] = set()

    def draw(self, gender: str) -> tuple[str, str]:
        while True:
            first = self.fake.first_name_female() if gender == "f" else self.fake.first_name_male()
            last = self.fake.last_name()
            names = (first.lower(), last.lower())
            if all(
                n.isalpha() and len(n) >= 3 and n not in BLOCKED_NAMES and n not in self.used
                for n in names
            ):
                self.used.update(names)
                return first, last


class GoldRecorder:
    """Collects unique ``(value, label, location)`` gold records."""

    def __init__(self) -> None:
        self.items: Gold = []

    def add(self, value: str, label: str, *locations: str) -> None:
        for location in locations:
            item = {"value": value, "label": label, "location": location}
            if item not in self.items:
                self.items.append(item)


def initials(name: str) -> str:
    return "".join(part[0] for part in name.split())


# --------------------------------------------------------------------------- #
# Raw-XML helpers for features python-docx has no API for
# --------------------------------------------------------------------------- #
def add_hyperlink(paragraph: Paragraph, url: str, text: str) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    paragraph._p.append(parse_xml(
        f'<w:hyperlink {nsdecls("w", "r")} r:id="{r_id}" w:history="1">'
        '<w:r><w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/></w:rPr>'
        f'<w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:hyperlink>'
    ))


def add_tracked_insertion(paragraph: Paragraph, text: str, author: str, revision_id: int) -> None:
    paragraph._p.append(parse_xml(
        f'<w:ins {nsdecls("w")} w:id="{revision_id}" w:author={quoteattr(author)} '
        f'w:date="{REVISION_DATE}"><w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:ins>'
    ))


def add_tracked_deletion(paragraph: Paragraph, text: str, author: str, revision_id: int) -> None:
    paragraph._p.append(parse_xml(
        f'<w:del {nsdecls("w")} w:id="{revision_id}" w:author={quoteattr(author)} '
        f'w:date="{REVISION_DATE}"><w:r><w:delText xml:space="preserve">{escape(text)}'
        "</w:delText></w:r></w:del>"
    ))


def add_text_box(paragraph: Paragraph, text: str) -> None:
    """A VML text box anchored in ``paragraph`` (its text is a separate paragraph)."""
    paragraph._p.append(parse_xml(
        f'<w:r {nsdecls("w")} xmlns:v="{VML_NS}"><w:pict>'
        '<v:rect style="width:260pt;height:40pt" strokecolor="#1F4E79" filled="f">'
        '<v:textbox><w:txbxContent><w:p><w:r>'
        f'<w:t xml:space="preserve">{escape(text)}</w:t>'
        "</w:r></w:p></w:txbxContent></v:textbox></v:rect></w:pict></w:r>"
    ))


def add_lines(document: Document, lines: list[str]) -> Paragraph:
    """One paragraph whose lines are separated by ``w:br`` line breaks."""
    paragraph = document.add_paragraph()
    for i, line in enumerate(lines):
        run = paragraph.add_run(line)
        if i < len(lines) - 1:
            run.add_break()
    return paragraph


def locality(fake: Faker) -> str:
    return f"{fake.city()} {fake.state_abbr()} {fake.postcode()}"


def birth_date(fake: Faker, start: int, end: int, fmt: str) -> str:
    return fake.date_between_dates(date(start, 1, 1), date(end, 12, 31)).strftime(fmt)


# --------------------------------------------------------------------------- #
# Documents
# --------------------------------------------------------------------------- #
def make_cv(path: Path, fake: Faker, rng: random.Random, people: People) -> Gold:
    """CV: header/footer PII, key-value table, split-run name, mailto link, comment."""
    first, last = people.draw("f")
    full = f"{first} {last}"
    reviewer = " ".join(people.draw("m"))
    recruiter = " ".join(people.draw("f"))
    email = f"{first}.{last}@example.com".lower()
    phone = make_mobile(rng)
    dob = birth_date(fake, 1968, 2000, "%d/%m/%Y")
    address = f"{fake.street_address()}, {locality(fake)}"
    tfn, medicare = make_tfn(rng), make_medicare(rng)
    bsb, account = make_bsb(rng), str(rng.randint(10_000_000, 99_999_999))
    employer = f"{fake.last_name()} Analytics Pty Ltd"
    assert tfn_is_valid(tfn) and medicare_is_valid(medicare)

    document = docx.Document()
    document.core_properties.author = recruiter
    document.core_properties.last_modified_by = recruiter
    document.core_properties.title = "Curriculum Vitae"
    section = document.sections[0]
    section.header.paragraphs[0].text = f"Curriculum Vitae of {full}"
    section.footer.paragraphs[0].text = f"{email} | {phone}"

    document.add_heading("Curriculum Vitae", level=0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph().add_run(f"Prepared for {full} (confidential)").bold = True
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in [
        ("Name", full), ("Date of birth", dob), ("Address", address), ("Phone", phone),
        ("Email", email), ("TFN", tfn), ("Medicare", medicare), ("BSB/Account", f"{bsb} / {account}"),
    ]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = key, value

    document.add_heading("Profile", level=1)
    document.add_paragraph(
        f"{full} is a data analyst at {employer}. Ms {last} led the migration of the "
        f"customer data warehouse in 2023, and {last}'s dashboards are used by the executive team."
    )
    contact = document.add_paragraph()
    contact.add_run(first).bold = True  # the name is split across two runs
    contact.add_run(f" {last}")
    contact.add_run(" can be reached at ")
    add_hyperlink(contact, f"mailto:{email}", email)
    contact.add_run(".")
    references = document.add_paragraph("References are available on request.")
    document.add_comment(
        references.runs,
        text=f"Confirm the referee details with {first} before sending.",
        author=reviewer,
        initials=initials(reviewer),
    )
    document.add_paragraph(
        "Available from May 2024; happy to work near the Sydney Opera House or remotely."
    )
    document.save(path)

    gold = GoldRecorder()
    gold.add(full, "PERSON", "header", "body", "table")
    gold.add(last, "PERSON", "body")
    gold.add(first, "PERSON", "comment")
    gold.add(email, "EMAIL", "footer", "table", "body")
    gold.add(phone, "PHONE", "footer", "table")
    gold.add(dob, "DATE", "table")
    gold.add(address, "ADDRESS", "table")
    gold.add(tfn, "ID_NUMBER", "table")
    gold.add(medicare, "ID_NUMBER", "table")
    gold.add(bsb, "FINANCIAL", "table")
    gold.add(account, "FINANCIAL", "table")
    gold.add(employer, "ORGANIZATION", "body")
    gold.add(reviewer, "PERSON", "comment_author")
    gold.add(recruiter, "PERSON", "metadata")
    return gold.items


def make_letter(path: Path, fake: Faker, rng: random.Random, people: People) -> Gold:
    """Referral letter: address blocks, tracked insertion/deletion, text box, hard negatives."""
    gp_first, gp_last = people.draw("m")
    gp = f"{gp_first} {gp_last}"
    specialist_first, specialist_last = people.draw("f")
    specialist = f"{specialist_first} {specialist_last}"
    patient_first, patient_last = people.draw("f")
    patient = f"{patient_first} {patient_last}"
    relative = " ".join(people.draw("m"))
    old_contact = " ".join(people.draw("m"))
    nurse = " ".join(people.draw("f"))  # author of the tracked changes
    manager = " ".join(people.draw("f"))  # core-properties author
    practice = f"{fake.last_name()} Street Medical Centre"
    clinic = f"{fake.city()} Cardiology Clinic"
    gp_street, gp_locality = fake.street_address(), locality(fake)
    clinic_street, clinic_locality = fake.street_address(), locality(fake)
    gp_phone = make_landline(rng)
    gp_email = f"{gp_first[0]}.{gp_last}@example.org".lower()
    dob = birth_date(fake, 1945, 1970, "%d %B %Y")
    medicare = make_medicare(rng)
    patient_phone = make_mobile(rng)
    assert medicare_is_valid(medicare)

    document = docx.Document()
    document.core_properties.author = manager
    document.core_properties.last_modified_by = manager
    document.core_properties.title = "Referral letter"
    add_lines(document, [f"Dr {gp}", practice, gp_street, gp_locality,
                         f"Phone: {gp_phone}", f"Email: {gp_email}"])
    document.add_paragraph("14 October 2025").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_lines(document, [f"Dr {specialist}", clinic, clinic_street, clinic_locality])
    document.add_paragraph(f"Dear Dr {specialist_last},")
    regarding = document.add_paragraph()
    regarding.add_run("Re: ").bold = True
    regarding.add_run(f"{patient}, DOB {dob}, Medicare {medicare}")
    body = document.add_paragraph(
        f"Thank you for seeing {patient_first}, a retired teacher with intermittent palpitations."
    )
    add_tracked_insertion(body, f" Her son, {relative}, will accompany her to appointments.", nurse, 901)
    add_tracked_deletion(body, f" Please copy all correspondence to {old_contact}.", nurse, 902)
    document.add_paragraph(
        "She was last reviewed in May 2024 after a fall near the Sydney Opera House. "
        "Please quote invoice number INV-20931 for the enclosed pathology results."
    )
    add_text_box(document.add_paragraph(), f"Urgent callback: {patient_phone}")
    add_lines(document, ["Yours sincerely,", f"Dr {gp}"])
    document.save(path)

    gold = GoldRecorder()
    gold.add(gp, "PERSON", "body")
    gold.add(practice, "ORGANIZATION", "body")
    gold.add(gp_street, "ADDRESS", "body")
    gold.add(gp_locality, "ADDRESS", "body")
    gold.add(gp_phone, "PHONE", "body")
    gold.add(gp_email, "EMAIL", "body")
    gold.add(specialist, "PERSON", "body")
    gold.add(specialist_last, "PERSON", "body")
    gold.add(clinic, "ORGANIZATION", "body")
    gold.add(clinic_street, "ADDRESS", "body")
    gold.add(clinic_locality, "ADDRESS", "body")
    gold.add(patient, "PERSON", "body")
    gold.add(patient_first, "PERSON", "body")
    gold.add(dob, "DATE", "body")
    gold.add(medicare, "ID_NUMBER", "body")
    gold.add(relative, "PERSON", "body")  # inside the (accepted) tracked insertion
    gold.add(old_contact, "PERSON", "deleted")
    gold.add(patient_phone, "PHONE", "textbox")
    gold.add(nurse, "PERSON", "metadata")
    gold.add(manager, "PERSON", "metadata")
    return gold.items


BUILDERS: dict[str, Callable[[Path, Faker, random.Random, People], Gold]] = {
    "sample_cv": make_cv,
    "sample_letter": make_letter,
}


def generate(out_dir: Path = DEFAULT_OUT, seed: int = 13) -> dict[str, Path]:
    """Write every sample ``.docx`` and its ``.gold.json`` into ``out_dir``."""
    out_dir.mkdir(parents=True, exist_ok=True)
    fake = Faker("en_AU")
    fake.seed_instance(seed)
    rng = random.Random(seed)
    people = People(fake)
    written: dict[str, Path] = {}
    for name, build in BUILDERS.items():
        path = out_dir / f"{name}.docx"
        gold = build(path, fake, rng, people)
        gold_path = out_dir / f"{name}.gold.json"
        gold_path.write_text(json.dumps(gold, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        written[name] = path
    return written


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate synthetic PII sample documents.")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help="output directory")
    parser.add_argument("--seed", type=int, default=13, help="random seed")
    args = parser.parse_args()
    for name, path in generate(args.out, args.seed).items():
        print(f"wrote {path} and {name}.gold.json")


if __name__ == "__main__":
    main()
