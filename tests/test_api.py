"""API tests with a stub model (no GPU or checkpoint needed)."""

from io import BytesIO

import pytest
from conftest import stub_detector
from docx import Document
from fastapi.testclient import TestClient

from redactner.api import create_app
from redactner.docx_engine import find_residuals

EMAIL = "jane.citizen@example.com"
PHONE = "0412 345 678"


def make_docx() -> bytes:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = f"Contact: {EMAIL}"
    p = doc.add_paragraph("Call ")
    p.add_run(PHONE).bold = True
    p.add_run(" any time.")
    doc.add_paragraph("No personal data in this sentence.")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def client():
    app = create_app(detector=stub_detector({EMAIL: "EMAIL", PHONE: "PHONE"}))
    with TestClient(app) as c:
        c.headers.update({"X-RedactNER-Token": app.state.token})
        yield c


def test_page_contains_token(client):
    html = client.get("/").text
    assert client.headers["X-RedactNER-Token"] in html
    assert "Content-Security-Policy" in client.get("/").headers


def test_token_and_origin_required(client):
    token = client.headers.pop("X-RedactNER-Token")
    try:
        assert client.get("/api/v1/config").status_code == 401
    finally:
        client.headers["X-RedactNER-Token"] = token
    assert client.get("/api/v1/config", headers={"Origin": "https://evil.example"}).status_code == 403
    assert client.get("/api/v1/config").status_code == 200


def test_upload_analyze_redact_download(client):
    up = client.post("/api/v1/documents", content=make_docx(),
                     headers={"Content-Type": "application/octet-stream", "X-Filename": "cv.docx"})
    assert up.status_code == 200, up.text
    doc_id = up.json()["doc_id"]
    res = client.post(f"/api/v1/documents/{doc_id}/analyze", json={})
    assert res.status_code == 200, res.text
    found = {(e["text"], e["label"]) for e in res.json()["entities"]}
    assert (EMAIL, "EMAIL") in found and (PHONE, "PHONE") in found
    red = client.post(f"/api/v1/documents/{doc_id}/redact", json={"mode": "label"})
    assert red.status_code == 200, red.text
    out = client.get(f"/api/v1/documents/{doc_id}/download")
    assert out.status_code == 200 and out.headers["cache-control"] == "no-store"
    assert find_residuals(out.content, [EMAIL, PHONE]) == []
    assert client.delete(f"/api/v1/documents/{doc_id}").json()["deleted"] is True


def test_rejects_garbage_and_validation_does_not_echo(client):
    bad = client.post("/api/v1/documents", content=b"not a docx")
    assert bad.status_code in (415, 422)
    up = client.post("/api/v1/documents", content=make_docx())
    doc_id = up.json()["doc_id"]
    r = client.post(f"/api/v1/documents/{doc_id}/analyze", json={"threshold": "SECRET-VALUE"})
    assert r.status_code == 422 and "SECRET-VALUE" not in r.text


def test_config_and_health_describe_the_model(client):
    cfg = client.get("/api/v1/config").json()
    assert "detectors" not in cfg and cfg["model"]["name"] == "stub-model"
    assert client.get("/api/v1/health").json()["status"] == "ok"


def test_missing_model_is_reported_clearly(tmp_path):
    from redactner.pipeline import ModelNotFoundError

    with pytest.raises(ModelNotFoundError):
        create_app(model_dir=tmp_path / "missing")
