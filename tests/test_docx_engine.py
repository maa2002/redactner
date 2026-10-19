"""Tests for the DOCX engine: safe open, segments, cross-run rewrite and leak checks."""

from __future__ import annotations

import bisect
import importlib.util
import io
import json
import re
import zipfile
from collections.abc import Callable
from pathlib import Path
from types import ModuleType
from xml.sax.saxutils import quoteattr

import docx
import pytest
from docx.document import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.shared import Twips

from redactner.docx_engine import (
    OBJECT_CHAR,
    W_NS,
    DocxDocument,
    DocxRefused,
    extract_text_raw,
    find_residuals,
)

ROOT = Path(__file__).resolve().parents[1]
Replacements = dict[str, list[tuple[int, int, str]]]


# --------------------------------------------------------------------------- #
# Fixtures and helpers
# --------------------------------------------------------------------------- #
def load_sample_script() -> ModuleType:
    spec = importlib.util.spec_from_file_location(
        "make_sample_docs", ROOT / "scripts" / "make_sample_docs.py"
    )
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="session")
def samples(tmp_path_factory: pytest.TempPathFactory) -> dict[str, tuple[bytes, list[dict]]]:
    """name -> (docx bytes, gold records), generated fresh into a temp dir."""
    out_dir = tmp_path_factory.mktemp("samples")
    paths = load_sample_script().generate(out_dir, seed=13)
    return {
        name: (
            path.read_bytes(),
            json.loads(path.with_name(f"{name}.gold.json").read_text(encoding="utf-8")),
        )
        for name, path in paths.items()
    }


def merge_spans(spans: list[tuple[int, int, str]]) -> list[tuple[int, int, str]]:
    merged: list[tuple[int, int, str]] = []
    for start, end, label in sorted(spans, key=lambda s: (s[0], -s[1])):
        if merged and start < merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], end), merged[-1][2])
        else:
            merged.append((start, end, label))
    return merged


def gold_replacements(doc: DocxDocument, gold: list[dict]) -> Replacements:
    """Locate gold values in document_text() and map them to segment-local spans."""
    text, starts = doc.document_text()
    spans: dict[str, list[tuple[int, int, str]]] = {}
    for value, label in {(g["value"], g["label"]) for g in gold}:
        for match in re.finditer(re.escape(value), text):
            i = bisect.bisect_right(starts, match.start()) - 1
            segment = doc.segments[i]
            start = match.start() - starts[i]
            end = min(match.end() - starts[i], len(segment.text))
            spans.setdefault(segment.id, []).append((start, end, f"[{label}]"))
    return {segment_id: merge_spans(found) for segment_id, found in spans.items()}


def gold_value(gold: list[dict], location: str, label: str = "PERSON") -> str:
    return next(g["value"] for g in gold if g["location"] == location and g["label"] == label)


def build_docx(build: Callable[[Document], object]) -> bytes:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def rezip(data: bytes, changes: dict[str, bytes]) -> bytes:
    """Copy a zip, replacing or adding the given entries."""
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename not in changes:
                target.writestr(info.filename, source.read(info.filename))
        for name, payload in changes.items():
            target.writestr(name, payload)
    return buffer.getvalue()


def read_part(data: bytes, name: str) -> str:
    return zipfile.ZipFile(io.BytesIO(data)).read(name).decode("utf-8")


def notes_xml(kind: str, text: str) -> bytes:
    """A footnotes/endnotes part with a separator note and one real note."""
    return (
        f'<w:{kind}s xmlns:w="{W_NS}">'
        f'<w:{kind} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{kind}>'
        f'<w:{kind} w:id="1"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:{kind}>'
        f"</w:{kind}s>"
    ).encode()


# A Word 2010+ text box: DrawingML in mc:Choice, the VML copy in mc:Fallback.
TEXT_BOX_RUN = (
    f'<w:r xmlns:w="{W_NS}" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml">'
    '<mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic>'
    "<a:graphicData><wps:wsp><wps:txbx><w:txbxContent>"
    "<w:p><w:r><w:t>Call Deanna Jennings</w:t></w:r></w:p>"
    "</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>"
    "</mc:Choice><mc:Fallback><w:pict><v:rect><v:textbox><w:txbxContent>"
    "<w:p><w:r><w:t>Call Deanna Jennings</w:t></w:r></w:p>"
    "</w:txbxContent></v:textbox></v:rect></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
)


# --------------------------------------------------------------------------- #
# Reading
# --------------------------------------------------------------------------- #
def test_segments_cover_every_location_in_reading_order(samples):
    data, gold = samples["sample_cv"]
    doc = DocxDocument.from_bytes(data)
    texts: dict[str, list[str]] = {}
    for segment in doc.segments:
        texts.setdefault(segment.location, []).append(segment.text)

    assert {"header", "body", "table", "footer", "comment"} <= texts.keys()
    assert any(gold_value(gold, "header") in t for t in texts["header"])
    assert any(gold_value(gold, "footer", "EMAIL") in t for t in texts["footer"])
    assert any(gold_value(gold, "comment") in t for t in texts["comment"])
    assert "Date of birth" in texts["table"]
    assert [s.location for s in doc.segments][0] == "header"
    assert [s.location for s in doc.segments][-2:] == ["footer", "comment"]
    assert [s.id for s in doc.segments] == [f"s{i}" for i in range(len(doc.segments))]
    assert doc.stats["tables"] == 1 and doc.stats["comments"] == 1
    assert doc.stats["headers"] == 1 and doc.stats["hyperlinks"] == 1


def test_document_text_offsets_align_with_segments(samples):
    doc = DocxDocument.from_bytes(samples["sample_cv"][0])
    text, starts = doc.document_text()
    assert len(starts) == len(doc.segments)
    for segment, start in zip(doc.segments, starts):
        assert text[start : start + len(segment.text)] == segment.text
    assert len(text) == starts[-1] + len(doc.segments[-1].text)


def test_tracked_changes_are_accepted(samples):
    data, gold = samples["sample_letter"]
    doc = DocxDocument.from_bytes(data)
    text, _ = doc.document_text()

    assert "will accompany her to appointments" in text  # insertion kept
    assert "copy all correspondence" not in text  # deletion dropped
    assert gold_value(gold, "deleted") not in text
    assert doc.stats["tracked_changes_accepted"] == 2
    assert [s.text for s in doc.segments if s.location == "textbox"] == [
        f"Urgent callback: {gold_value(gold, 'textbox', 'PHONE')}"
    ]
    assert all(s.text.strip(OBJECT_CHAR) for s in doc.segments)  # no image-only segments


def test_notes_follow_the_body_and_precede_footers(samples):
    data = samples["sample_cv"][0]
    overrides = "".join(
        f'<Override PartName="/word/{kind}s.xml" ContentType="application/'
        f'vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml"/>'
        for kind in ("footnote", "endnote")
    )
    types = read_part(data, "[Content_Types].xml").replace("</Types>", overrides + "</Types>")
    data = rezip(data, {
        "[Content_Types].xml": types.encode(),
        "word/footnotes.xml": notes_xml("footnote", "A footnote."),
        "word/endnotes.xml": notes_xml("endnote", "An endnote."),
    })
    doc = DocxDocument.from_bytes(data)
    order = list(dict.fromkeys(s.location for s in doc.segments if s.location != "table"))
    assert order == ["header", "body", "footnote", "endnote", "footer", "comment"]
    assert doc.stats["footnotes"] == 1 and doc.stats["endnotes"] == 1


def test_text_box_choice_and_fallback_copies_are_both_redacted():
    doc = DocxDocument.from_bytes(
        build_docx(lambda d: d.add_paragraph("Anchor text ")._p.append(parse_xml(TEXT_BOX_RUN)))
    )
    boxes = [s for s in doc.segments if s.location == "textbox"]
    assert [s.text for s in boxes] == ["Call Deanna Jennings"] * 2
    assert doc.segments[0].text == f"Anchor text {OBJECT_CHAR}"  # one object char, Choice branch only

    out, report = doc.redact({boxes[0].id: [(5, 20, "[PERSON]")]})
    assert report["residual_fixes"] == 1  # the mc:Fallback copy
    after = [s.text for s in DocxDocument.from_bytes(out).segments if s.location == "textbox"]
    assert after == ["Call [PERSON]", "Call [REDACTED]"]
    assert find_residuals(out, ["Deanna Jennings"]) == []


# --------------------------------------------------------------------------- #
# Layout hints
# --------------------------------------------------------------------------- #
def test_layout_reads_style_alignment_and_numbering():
    def build(document: Document) -> None:
        document.add_heading("Quarterly report", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph("Justified body.").alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        document.add_paragraph("Right aligned.").alignment = WD_ALIGN_PARAGRAPH.RIGHT
        item = document.add_paragraph("First item", style="List Paragraph")
        numbering = item._p.get_or_add_pPr().get_or_add_numPr()
        numbering.get_or_add_ilvl().val = 0
        numbering.get_or_add_numId().val = 1
        document.add_paragraph("Plain.")

    segments = {s.text: s for s in DocxDocument.from_bytes(build_docx(build)).segments}
    heading = segments["Quarterly report"]
    assert (heading.style, heading.align, heading.list) == ("Heading1", "center", False)
    assert segments["Justified body."].align == "both"
    assert segments["Right aligned."].align == "right"
    assert (segments["First item"].style, segments["First item"].list) == ("ListParagraph", True)
    plain = segments["Plain."]
    assert (plain.style, plain.align, plain.list, plain.table) == (None, None, False, None)
    assert (plain.marks, plain.hidden, plain.author) == ([], False, None)


def test_layout_marks_tables_and_comment_authors(samples):
    data, gold = samples["sample_cv"]
    doc = DocxDocument.from_bytes(data)
    first, email = gold_value(gold, "comment"), gold_value(gold, "body", "EMAIL")

    contact = next(s for s in doc.segments if "can be reached" in s.text)
    start = contact.text.index(email)
    assert contact.marks == [
        {"s": 0, "e": len(first), "b": True, "i": False, "u": False},  # bold first-name run
        {"s": start, "e": start + len(email), "b": False, "i": False, "u": True},  # link run
    ]
    prepared = next(s for s in doc.segments if s.text.startswith("Prepared for"))
    assert prepared.marks == [{"s": 0, "e": len(prepared.text), "b": True, "i": False, "u": False}]
    title = next(s for s in doc.segments if s.text == "Curriculum Vitae")
    assert (title.style, title.align) == ("Title", "center")

    label = next(s for s in doc.segments if s.text == "Date of birth")
    value = next(s for s in doc.segments if s.text == gold_value(gold, "table", "DATE"))
    assert label.table == {"t": 0, "r": 1, "c": 0, "rows": 8, "cols": 2}
    assert value.table == {"t": 0, "r": 1, "c": 1, "rows": 8, "cols": 2}
    assert all(s.table is None for s in doc.segments if s.location != "table")

    (comment,) = [s for s in doc.segments if s.location == "comment"]
    assert comment.author == gold_value(gold, "comment_author")
    assert all(s.author is None for s in doc.segments if s is not comment)
    assert json.loads(json.dumps([s.to_dict() for s in doc.segments]))[0]["id"] == "s0"


def test_nested_tables_get_their_own_index():
    def build(document: Document) -> None:
        outer = document.add_table(rows=2, cols=2)
        outer.cell(1, 0).text = "Outer cell"
        inner = outer.cell(1, 1).add_table(rows=1, cols=3)
        for c, cell in enumerate(inner.rows[0].cells):
            cell.text = f"Inner {c}"

    segments = {s.text: s for s in DocxDocument.from_bytes(build_docx(build)).segments}
    assert segments["Outer cell"].table == {"t": 0, "r": 1, "c": 0, "rows": 2, "cols": 2}
    assert segments["Inner 2"].table == {"t": 1, "r": 0, "c": 2, "rows": 1, "cols": 3}


def test_text_box_fallback_copy_is_hidden():
    doc = DocxDocument.from_bytes(
        build_docx(lambda d: d.add_paragraph("Anchor ")._p.append(parse_xml(TEXT_BOX_RUN)))
    )
    boxes = [s for s in doc.segments if s.location == "textbox"]
    assert [s.hidden for s in boxes] == [False, True]  # mc:Choice visible, mc:Fallback hidden
    assert not doc.segments[0].hidden


def test_sample_identifiers_have_valid_checksums(samples):
    script = load_sample_script()
    ids = [g["value"] for g in samples["sample_cv"][1] if g["label"] == "ID_NUMBER"]
    tfn, medicare = ids
    assert script.tfn_is_valid(tfn) and script.medicare_is_valid(medicare)


# --------------------------------------------------------------------------- #
# Redaction
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("name", ["sample_cv", "sample_letter"])
def test_gold_redaction_leaves_no_residuals(samples, name):
    data, gold = samples[name]
    values = sorted({g["value"] for g in gold})
    assert find_residuals(data, values)  # sanity: the PII really is in the input

    doc = DocxDocument.from_bytes(data)
    out, report = doc.redact(gold_replacements(doc, gold))

    # Includes values only present in metadata, comment authors, hyperlink
    # targets and deleted tracked changes -- none were passed as replacements.
    assert find_residuals(out, values) == []
    raw = extract_text_raw(out)
    for value in values:
        assert not re.search(rf"(?<!\w){re.escape(value)}(?!\w)", raw, re.IGNORECASE), value
    assert docx.Document(io.BytesIO(out)).paragraphs  # opens in python-docx
    assert len(DocxDocument.from_bytes(out).segments) == len(doc.segments)
    assert report["replacements_applied"] == sum(report["by_location"].values()) > 0


def test_untouched_text_keeps_formatting_and_split_names_are_replaced(samples):
    data, gold = samples["sample_cv"]
    doc = DocxDocument.from_bytes(data)
    out, report = doc.redact(gold_replacements(doc, gold))
    paragraphs = docx.Document(io.BytesIO(out)).paragraphs

    prepared = next(p for p in paragraphs if p.text.startswith("Prepared for"))
    assert prepared.text == "Prepared for [PERSON] (confidential)"
    assert prepared.runs and all(run.bold for run in prepared.runs)

    contact = next(p for p in paragraphs if "can be reached" in p.text)
    assert contact.text == "[PERSON] can be reached at [EMAIL]."  # bold first + plain surname
    assert contact.runs[0].bold and contact.runs[0].text == "[PERSON]"
    assert not contact.runs[1].bold and contact.runs[1].text == ""

    assert report["hyperlinks_unlinked"] == 1
    assert "<w:hyperlink" not in read_part(out, "word/document.xml")
    rels = read_part(out, "word/_rels/document.xml.rels")
    assert "https://redacted.invalid/" in rels and "mailto:" not in rels


def test_hard_negatives_survive(samples):
    data, gold = samples["sample_letter"]
    doc = DocxDocument.from_bytes(data)
    out, _ = doc.redact(gold_replacements(doc, gold))
    text, _ = DocxDocument.from_bytes(out).document_text()
    for phrase in ("May 2024", "Sydney Opera House", "INV-20931", "14 October 2025"):
        assert phrase in text


def test_metadata_comment_authors_and_thumbnail_are_scrubbed(samples):
    data, gold = samples["sample_cv"]
    out, report = DocxDocument.from_bytes(data).redact({})
    reopened = docx.Document(io.BytesIO(out))

    assert reopened.core_properties.author == ""
    assert reopened.core_properties.last_modified_by == ""
    assert [comment.author for comment in reopened.comments] == ["Reviewer 1"]
    assert report["comment_authors_anonymised"] == 1
    assert report["thumbnail_removed"]
    names = zipfile.ZipFile(io.BytesIO(out)).namelist()
    assert not any(name.startswith("docProps/thumbnail") for name in names)
    assert "thumbnail" not in read_part(out, "_rels/.rels")

    kept, kept_report = DocxDocument.from_bytes(data).redact({}, scrub_metadata=False)
    assert docx.Document(io.BytesIO(kept)).core_properties.author == gold_value(gold, "metadata")
    assert not kept_report["thumbnail_removed"]


def test_people_part_shares_the_comment_author_mapping(samples):
    data, gold = samples["sample_cv"]
    people = (
        '<w15:people xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml">'
        f"<w15:person w15:author={quoteattr(gold_value(gold, 'comment_author'))}>"
        '<w15:presenceInfo w15:providerId="AD" w15:userId="S::someone@example.com::42"/>'
        "</w15:person></w15:people>"
    )
    out, report = DocxDocument.from_bytes(rezip(data, {"word/people.xml": people.encode()})).redact({})
    people_out = read_part(out, "word/people.xml")
    assert 'w15:author="Reviewer 1"' in people_out and "presenceInfo" not in people_out
    assert report["comment_authors_anonymised"] == 1


def test_residual_verifier_catches_copies_the_caller_missed(samples):
    data, gold = samples["sample_cv"]
    doc = DocxDocument.from_bytes(data)
    full, surname = gold_value(gold, "header"), gold_value(gold, "body").split()[-1]
    header = doc.segments[0]
    start = header.text.index(full)

    out, report = doc.redact(
        {header.id: [(start, start + len(full), "[PERSON]")]}, redacted_values=[surname]
    )
    assert report["replacements_applied"] == 1
    assert report["residual_fixes"] > 0  # table, body, split runs, possessive...
    assert find_residuals(out, [full, surname]) == []


def test_keep_values_survive_the_safety_net():
    def build(document: Document) -> None:
        document.add_paragraph("Primary: jane@example.com, 0412 345 678")
        document.add_paragraph("Backup: jane@example.com, 0412-345-678 (Jane)")

    doc = DocxDocument.from_bytes(build_docx(build))
    second = doc.segments[1]
    email, phone = second.text.index("jane@"), second.text.index("0412")
    out, _ = doc.redact(
        {second.id: [(email, email + 16, "[EMAIL]"), (phone, phone + 12, "[PHONE]")]},
        redacted_values=["Jane"],  # also redacted, but must not eat into the kept e-mail
        keep_values=["JANE@example.com", "0412345678"],  # case / separators ignored
    )
    assert [s.text for s in DocxDocument.from_bytes(out).segments] == [
        "Primary: jane@example.com, 0412 345 678",
        "Backup: [EMAIL], [PHONE] ([REDACTED])",
    ]
    assert extract_text_raw(out).count("jane@example.com") == 1


def test_redact_twice_starts_from_the_original(samples):
    data, gold = samples["sample_cv"]
    doc = DocxDocument.from_bytes(data)
    full = gold_value(gold, "header")
    header = doc.segments[0]
    start = header.text.index(full)
    footer = next(s for s in doc.segments if s.location == "footer")

    first, _ = doc.redact({header.id: [(start, start + len(full), "[PERSON]")]})
    second, _ = doc.redact({footer.id: [(0, footer.text.index(" |"), "[EMAIL]")]})

    first_text = DocxDocument.from_bytes(first).document_text()[0]
    second_text = DocxDocument.from_bytes(second).document_text()[0]
    assert full not in first_text
    assert full in second_text and "[EMAIL] |" in second_text
    assert doc.segments[0].text == header.text  # loaded state untouched


def test_non_text_pieces_inside_a_span_are_removed():
    def build(document: Document) -> None:
        run = document.add_paragraph().add_run("Call 0412")
        for digits in ("345", "678"):
            run._r.append(OxmlElement("w:noBreakHyphen"))
            text = OxmlElement("w:t")
            text.text = digits
            run._r.append(text)
        run.add_tab()
        run.add_text("today")

    doc = DocxDocument.from_bytes(build_docx(build))
    (segment,) = doc.segments
    assert segment.text == "Call 0412-345-678\ttoday"

    out, _ = doc.redact({segment.id: [(5, 17, "[PHONE]")]})
    assert DocxDocument.from_bytes(out).segments[0].text == "Call [PHONE]\ttoday"
    assert "noBreakHyphen" not in read_part(out, "word/document.xml")

    # A span covering only a tab still receives its replacement text.
    out, _ = doc.redact({segment.id: [(17, 18, "[GAP]")]})
    assert DocxDocument.from_bytes(out).segments[0].text == "Call 0412-345-678[GAP]today"


def test_values_split_mid_word_across_runs_are_found_and_fixed():
    def build(document: Document) -> None:
        paragraph = document.add_paragraph("Signed: Deanna Jen")
        paragraph.add_run("nin").italic = True
        paragraph.add_run("gs, 3 May")

    data = build_docx(build)
    assert [h["where"] for h in find_residuals(data, ["deanna jennings"])] == ["text"]
    out, report = DocxDocument.from_bytes(data).redact({}, redacted_values=["Deanna Jennings"])
    assert report["residual_fixes"] == 1
    assert DocxDocument.from_bytes(out).segments[0].text == "Signed: [REDACTED], 3 May"
    assert find_residuals(out, ["Deanna Jennings"]) == []


def test_safety_net_never_rewrites_schema_tokens_or_numbers_in_markup():
    def build(document: Document) -> None:
        paragraph = document.add_paragraph("Contact Ms Page at 2000 Main St.")
        paragraph.paragraph_format.space_before = Twips(2000)
        paragraph.runs[0].add_break(WD_BREAK.PAGE)

    out, _ = DocxDocument.from_bytes(build_docx(build)).redact({}, redacted_values=["Page", "2000"])
    xml = read_part(out, "word/document.xml")
    assert 'w:type="page"' in xml and 'w:before="2000"' in xml  # enum and twips intact
    assert DocxDocument.from_bytes(out).segments[0].text == "Contact Ms [REDACTED] at [REDACTED] Main St.\n"
    assert find_residuals(out, ["Page", "2000"]) == []
    assert docx.Document(io.BytesIO(out)).paragraphs


def test_matching_rules_word_boundaries_digits_and_short_values():
    data = build_docx(lambda d: (d.add_paragraph("Annual report for Ann."),
                                 d.add_paragraph("Mobile 0412-345-678 (cell).")))
    hits = find_residuals(data, ["Ann", "0412 345 678", "Mo"])
    assert sorted((h["value"], h["part"], h["where"]) for h in hits) == [
        ("0412 345 678", "word/document.xml", "text"),
        ("Ann", "word/document.xml", "text"),
    ]


# --------------------------------------------------------------------------- #
# Refusals
# --------------------------------------------------------------------------- #
def test_refuses_encrypted_or_legacy_files():
    with pytest.raises(DocxRefused, match="encrypted"):
        DocxDocument.from_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(512))


def test_refuses_non_zip_input():
    with pytest.raises(DocxRefused, match="zip"):
        DocxDocument.from_bytes(b"plain text, not a document")


def test_refuses_macro_enabled_package(samples):
    data = rezip(samples["sample_cv"][0], {"word/vbaProject.bin": b"\x00" * 64})
    with pytest.raises(DocxRefused, match="macro"):
        DocxDocument.from_bytes(data)


@pytest.mark.parametrize(
    ("old", "new", "message"),
    [
        (b"?>", b'?><!DOCTYPE w:document [<!ENTITY x "boom">]>', "DOCTYPE"),
        (W_NS.encode(), b"http://purl.oclc.org/ooxml/wordprocessingml/main", "Strict"),
        (b"<w:sectPr", b'<w:altChunk r:id="rId99"/><w:sectPr', "altChunk"),
    ],
)
def test_refuses_unsafe_document_xml(samples, old, new, message):
    data = samples["sample_cv"][0]
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml")
    assert old in xml
    with pytest.raises(DocxRefused, match=message):
        DocxDocument.from_bytes(rezip(data, {"word/document.xml": xml.replace(old, new, 1)}))


@pytest.mark.parametrize(
    "limits", [{"max_bytes": 1_000}, {"max_entries": 5}, {"max_uncompressed": 20_000}]
)
def test_enforces_size_limits(samples, limits):
    with pytest.raises(DocxRefused):
        DocxDocument.from_bytes(samples["sample_cv"][0], **limits)
